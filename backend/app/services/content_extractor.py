"""Content extraction and normalisation (Phase 3).

Supports webpages (URL or raw HTML), YouTube videos (title, description, transcript),
plain text and uploaded documents (PDF, DOCX, TXT/MD).
"""
import io
import ipaddress
import json
import os
import re
import socket
import time
import urllib.parse
from concurrent.futures import ThreadPoolExecutor, wait
from pathlib import Path
from typing import Any, Dict, List, Optional

import requests
from bs4 import BeautifulSoup
from youtube_transcript_api import YouTubeTranscriptApi

MAX_WORDS = 20_000
MAX_FETCH_BYTES = 3_000_000
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
SUPPORTED_DOC_TYPES = {".pdf", ".docx", ".txt", ".md", ".csv"}
_UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) FocusAI/1.0"}
_BLOCK_TAGS = ["p", "h1", "h2", "h3", "h4", "li", "blockquote"]


# --------------------------------------------------------------------------- helpers
def clean_text(text: str) -> str:
    """Collapse whitespace/newlines and strip control characters."""
    if not text:
        return ""
    text = re.sub(r"[\x00-\x08\x0b-\x1f\x7f]", " ", text)
    text = re.sub(r"[\r\n\t]+", " ", text)
    text = re.sub(r"\s{2,}", " ", text)
    return text.strip()


def _limit_words(text: str) -> str:
    words = text.split()
    return text if len(words) <= MAX_WORDS else " ".join(words[:MAX_WORDS])


def _result(source_type: str, title: str, text: str, url: Optional[str] = None, **metadata: Any) -> Dict[str, Any]:
    cleaned = _limit_words(clean_text(text))
    return {
        "source_type": source_type,
        "url": url,
        "title": title,
        "cleaned_text": cleaned,
        "word_count": len(cleaned.split()),
        "char_count": len(cleaned),
        "metadata": metadata,
    }


def _assert_public_url(url: str) -> None:
    """Block SSRF: only http(s) to public IP addresses (set ALLOW_PRIVATE_URLS=1 to disable)."""
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.hostname:
        raise ValueError("Only http(s) URLs are supported.")
    if os.getenv("ALLOW_PRIVATE_URLS") == "1":
        return
    for info in socket.getaddrinfo(parsed.hostname, None):
        ip = ipaddress.ip_address(info[4][0])
        if ip.is_private or ip.is_loopback or ip.is_link_local or ip.is_reserved or ip.is_multicast:
            raise ValueError("URLs that resolve to private or local addresses are not allowed.")


def fetch_url(url: str, timeout: float = 8.0, max_redirects: int = 3) -> str:
    """GET a URL safely, validating every redirect hop. Returns response text."""
    for _ in range(max_redirects + 1):
        _assert_public_url(url)
        resp = requests.get(url, headers=_UA, timeout=timeout, allow_redirects=False)
        if resp.status_code in (301, 302, 303, 307, 308) and resp.headers.get("Location"):
            url = urllib.parse.urljoin(url, resp.headers["Location"])
            continue
        resp.raise_for_status()
        return resp.text[:MAX_FETCH_BYTES]
    raise ValueError("Too many redirects.")


# --------------------------------------------------------------------------- YouTube
def extract_youtube_id(url: str) -> Optional[str]:
    """Extract the video ID from watch / embed / shorts / youtu.be URLs."""
    if not url:
        return None
    parsed = urllib.parse.urlparse(url)
    host = (parsed.hostname or "").lower()
    if host in ("www.youtube.com", "youtube.com", "m.youtube.com", "music.youtube.com"):
        if parsed.path == "/watch":
            return urllib.parse.parse_qs(parsed.query).get("v", [None])[0]
        parts = [p for p in parsed.path.split("/") if p]
        if len(parts) >= 2 and parts[0] in ("embed", "v", "shorts", "live"):
            return parts[1]
    elif host == "youtu.be":
        return parsed.path.lstrip("/") or None
    return None


# ---- text normalisation for video text -------------------------------------------------------
_URL_RE = re.compile(r"https?://\S+|www\.\S+")
_EMAIL_RE = re.compile(r"\S+@\S+\.\S+")
_TIMESTAMP_RE = re.compile(r"\b\d{1,2}:\d{2}(?::\d{2})?\b")
_CAPTION_NOISE_RE = re.compile(r"\[[^\]]{1,30}\]|>>+")  # [Music], [Applause], >> speaker markers
MAX_DESCRIPTION_CHARS = 2000
MAX_TRANSCRIPT_WORDS = 6000
MAX_TAGS = 15
YT_FETCH_TIMEOUT = float(os.getenv("YT_FETCH_TIMEOUT_SECONDS", 9))
YT_CACHE_TTL = 3600


def normalize_video_text(text: str, limit_chars: Optional[int] = None) -> str:
    """Clean descriptions/transcripts: drop URLs, emails, timestamps and caption noise, collapse whitespace."""
    if not text:
        return ""
    text = _URL_RE.sub(" ", text)
    text = _EMAIL_RE.sub(" ", text)
    text = _TIMESTAMP_RE.sub(" ", text)
    text = _CAPTION_NOISE_RE.sub(" ", text)
    text = re.sub(r"([!?.\-_=*~#])\1{2,}", r"\1", text)  # "!!!!!!" / "------"
    text = clean_text(text)
    return text[:limit_chars] if limit_chars else text


def is_generic_video_title(title: str) -> bool:
    """YouTube's SPA often still shows the site name (or '(3) YouTube') while the new video loads."""
    t = re.sub(r"^\(\d+\)\s*", "", (title or "").strip()).strip(" -|").lower()
    return t in ("", "youtube", "youtube video")


# ---- per-video cache so re-visits / re-analysis do not refetch ---------------------------------
_YT_CACHE: Dict[str, Dict[str, Any]] = {}


def clear_youtube_cache() -> None:
    _YT_CACHE.clear()


def _cache_get(video_id: str, part: str) -> Any:
    entry = _YT_CACHE.get(video_id, {})
    if part in entry and time.time() - entry[part][0] < YT_CACHE_TTL:
        return entry[part][1]
    return None


def _cache_put(video_id: str, part: str, value: Any) -> None:
    if len(_YT_CACHE) > 256:
        _YT_CACHE.clear()
    _YT_CACHE.setdefault(video_id, {})[part] = (time.time(), value)


# ---- the three network fetches (each best-effort, run in parallel with a shared deadline) ------
def _fetch_transcript(video_id: str) -> str:
    """Transcript text via youtube-transcript-api >= 1.0 (``fetch``), preferring English."""
    api = YouTubeTranscriptApi()
    try:
        fetched = api.fetch(video_id, languages=["en", "en-US", "en-GB"])
    except Exception:
        # No English track: take whichever transcript exists.
        fetched = next(iter(api.list(video_id))).fetch()
    return " ".join(snippet.text for snippet in fetched)


def _fetch_oembed_title(video_id: str) -> str:
    res = requests.get(
        "https://www.youtube.com/oembed",
        params={"url": f"https://www.youtube.com/watch?v={video_id}", "format": "json"},
        timeout=5,
    )
    return res.json().get("title", "") if res.status_code == 200 else ""


def _parse_player_response(html: str) -> Optional[Dict[str, Any]]:
    m = re.search(r"ytInitialPlayerResponse\s*=\s*", html)
    if not m:
        return None
    try:
        obj, _ = json.JSONDecoder().raw_decode(html[m.end():])
        return obj if isinstance(obj, dict) else None
    except ValueError:
        return None


def _youtube_page_meta(video_id: str) -> Dict[str, Any]:
    """Best-effort title / full description / tags / channel / category from the public watch page.

    Reads the player response embedded in the page (authoritative for *this* video id) and falls
    back to Open Graph tags. Consent walls or markup changes simply yield an empty result.
    """
    meta: Dict[str, Any] = {"title": "", "description": "", "keywords": [], "channel": "", "category": ""}
    try:
        html = fetch_url(f"https://www.youtube.com/watch?v={video_id}", timeout=5)
        player = _parse_player_response(html) or {}
        details = player.get("videoDetails") or {}
        micro = (player.get("microformat") or {}).get("playerMicroformatRenderer") or {}
        meta["title"] = clean_text(details.get("title", ""))
        meta["description"] = details.get("shortDescription", "") or ""
        meta["keywords"] = [str(k) for k in (details.get("keywords") or [])][:MAX_TAGS]
        meta["channel"] = clean_text(details.get("author", ""))
        meta["category"] = clean_text(micro.get("category", ""))
        if not meta["title"] or not meta["description"]:
            soup = BeautifulSoup(html, "lxml")
            for key, prop in (("title", "og:title"), ("description", "og:description")):
                tag = soup.find("meta", attrs={"property": prop})
                if not meta[key] and tag and tag.get("content"):
                    meta[key] = clean_text(tag["content"])
    except Exception:
        pass
    return meta


def _run_parallel(jobs: Dict[str, Any], timeout: float) -> Dict[str, Any]:
    """Run {name: callable} concurrently. Returns {name: ("ok", value) | ("error", exc) | ("timeout", None)}."""
    results: Dict[str, Any] = {}
    if not jobs:
        return results
    pool = ThreadPoolExecutor(max_workers=len(jobs))
    futures = {name: pool.submit(fn) for name, fn in jobs.items()}
    done, _pending = wait(futures.values(), timeout=timeout)
    for name, fut in futures.items():
        if fut in done:
            exc = fut.exception()
            results[name] = ("error", exc) if exc else ("ok", fut.result())
        else:
            results[name] = ("timeout", None)
    pool.shutdown(wait=False, cancel_futures=True)  # never block the request on a slow fetch
    return results


def extract_youtube_content(
    video_id: str,
    url: str,
    client_title: str = "",
    client_description: str = "",
    client_channel: str = "",
    client_keywords: Optional[List[str]] = None,
) -> Dict[str, Any]:
    """Title + description + metadata + transcript for ONE video. Every part is best-effort.

    ``client_*`` values come from the extension (what the user's page actually shows). Server-side
    values are authoritative for the video id; client values fill the gaps. A missing transcript
    (disabled captions, blocked IP, timeout...) never fails the extraction: the caller still gets
    title/description/metadata and ``metadata["transcript_status"]`` says why.
    """
    jobs: Dict[str, Any] = {}
    cached = {p: _cache_get(video_id, p) for p in ("title", "meta", "transcript")}
    if cached["title"] is None:
        jobs["title"] = lambda: _fetch_oembed_title(video_id)
    if cached["meta"] is None:
        jobs["meta"] = lambda: _youtube_page_meta(video_id)
    if cached["transcript"] is None:
        jobs["transcript"] = lambda: _fetch_transcript(video_id)
    fetched = _run_parallel(jobs, YT_FETCH_TIMEOUT)

    def value(part: str, default: Any) -> Any:
        if cached[part] is not None:
            return cached[part]
        status, val = fetched.get(part, ("ok", default))
        if status == "ok" and val:
            _cache_put(video_id, part, val)
            return val
        return default

    oembed_title = value("title", "") or ""
    page = value("meta", {}) or {}
    transcript_raw = value("transcript", "") or ""

    transcript_status = "ok"
    if not transcript_raw:
        status, exc = fetched.get("transcript", ("error", None))
        transcript_status = "timeout" if status == "timeout" else f"unavailable:{type(exc).__name__}" if exc else "unavailable"

    # Title: oEmbed is authoritative for this video id; then the player response; then the client's
    # page title (but never the generic site name YouTube shows while a new video loads).
    title = next((t for t in (oembed_title, page.get("title", ""), client_title) if t and not is_generic_video_title(t)), "")

    description = max(
        (normalize_video_text(d, MAX_DESCRIPTION_CHARS) for d in (page.get("description", ""), client_description)),
        key=len,
    )
    tags: List[str] = []
    for tag in list(page.get("keywords", [])) + list(client_keywords or []):
        tag = clean_text(str(tag)).lstrip("#")
        if tag and tag.lower() not in (t.lower() for t in tags):
            tags.append(tag)
    tags = tags[:MAX_TAGS]
    channel = page.get("channel", "") or clean_text(client_channel)
    category = page.get("category", "")
    metadata_text = " | ".join(
        p for p in (
            f"Channel: {channel}" if channel else "",
            f"Category: {category}" if category else "",
            f"Tags: {', '.join(tags)}" if tags else "",
        ) if p
    )

    transcript = " ".join(normalize_video_text(transcript_raw).split()[:MAX_TRANSCRIPT_WORDS])
    out = _result(
        "youtube", title or f"YouTube Video ({video_id})",
        ". ".join(p for p in (title, description, metadata_text, transcript) if p), url,
        video_id=video_id, title_known=bool(title), has_description=bool(description),
        has_transcript=bool(transcript), transcript_status=transcript_status,
        channel=channel, category=category, keywords=tags,
    )
    # Separate parts for the relevance pipeline (title / description / metadata / transcript are
    # weighted and presented to the LLM as distinct, labelled fields).
    out["parts"] = {"title": title, "description": description, "metadata_text": metadata_text, "transcript": transcript}
    return out


# --------------------------------------------------------------------------- webpages
def extract_webpage_content(url: Optional[str] = None, html: Optional[str] = None) -> Dict[str, Any]:
    """Main text + title from a URL or raw HTML. Returns ``error`` (and empty text) if fetching fails."""
    if not html and url:
        try:
            html = fetch_url(url)
        except Exception as exc:
            out = _result("webpage", "Untitled Webpage", "", url)
            out["error"] = str(exc)
            return out
    if not html:
        return _result("webpage", "Untitled Webpage", "", url)

    soup = BeautifulSoup(html, "lxml")
    title = "Untitled Webpage"
    if soup.title and soup.title.string and soup.title.string.strip():
        title = clean_text(soup.title.string)
    elif (og := soup.find("meta", attrs={"property": "og:title"})) and og.get("content"):
        title = clean_text(og["content"])

    for tag in soup(["script", "style", "nav", "header", "footer", "aside", "form", "noscript", "svg", "iframe"]):
        tag.decompose()

    area = (
        soup.find("article")
        or soup.find("main")
        or soup.find("div", {"id": re.compile(r"content|main", re.I)})
        or soup.body
        or soup
    )

    blocks, seen = [], set()
    for el in area.find_all(_BLOCK_TAGS):
        if el.name in ("li", "blockquote") and el.find(["p", "li", "blockquote"]):
            continue  # its nested blocks are visited on their own
        text = clean_text(el.get_text(" "))
        if not text or text in seen or (el.name in ("p", "li", "blockquote") and len(text.split()) < 3):
            continue
        seen.add(text)
        blocks.append(text if text[-1] in '.!?:;"\u201d)' else text + ".")  # keeps sentence boundaries
    body = " ".join(blocks)
    if not blocks:  # unusual markup (no paragraphs/headings): fall back to everything visible
        body = clean_text(area.get_text(" "))

    domain = urllib.parse.urlparse(url).netloc if url else None
    return _result("webpage", title, body, url, domain=domain)


# --------------------------------------------------------------------------- text & documents
def process_text_content(text: str, title: Optional[str] = None) -> Dict[str, Any]:
    cleaned = clean_text(text)
    doc_title = title or (cleaned[:50] + "..." if len(cleaned) > 50 else cleaned or "Document Text")
    return _result("document", doc_title, cleaned)


def extract_document(filename: str, data: bytes) -> Dict[str, Any]:
    """Extract text from an uploaded PDF, DOCX or plain-text file."""
    ext = Path(filename or "").suffix.lower()
    if ext not in SUPPORTED_DOC_TYPES:
        raise ValueError(f"Unsupported file type '{ext or 'unknown'}'. Supported: {', '.join(sorted(SUPPORTED_DOC_TYPES))}")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ValueError("File is too large (10 MB maximum).")

    pages = None
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = len(reader.pages)
        text = "\n".join((page.extract_text() or "") for page in reader.pages[:300])
    elif ext == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            parts += [cell.text for row in table.rows for cell in row.cells]
        text = "\n".join(parts)
    else:
        text = data.decode("utf-8", errors="replace")

    out = _result("document", Path(filename).stem, text, filename=filename, file_type=ext.lstrip("."), pages=pages)
    if not out["cleaned_text"]:
        out["error"] = "No extractable text found (scanned PDFs need OCR, which is not supported yet)."
    return out


def extract_and_normalize(
    url: Optional[str] = None, html: Optional[str] = None, text: Optional[str] = None, title: Optional[str] = None
) -> Dict[str, Any]:
    """Unified entry point to extract, clean and normalise online or user content."""
    if url:
        yt_id = extract_youtube_id(url)
        if yt_id:
            return extract_youtube_content(yt_id, url)
        if text and not html:
            out = process_text_content(text, title)
            out.update(source_type="webpage", url=url)
            return out
        return extract_webpage_content(url=url, html=html)
    if html:
        return extract_webpage_content(html=html)
    if text:
        return process_text_content(text, title=title)
    return _result("unknown", "Empty Payload", "")
